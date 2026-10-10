"""Importe le programme des ateliers depuis le planning Excel (lecture seule) vers data/seed/workshops.json."""
import json, re, sys, warnings
from pathlib import Path
import openpyxl

warnings.filterwarnings('ignore')
SRC = Path(sys.argv[1] if len(sys.argv) > 1 else
           '/Users/neiba/myspace/clients/cybastion/Planning_ateliers_cadrage_mis_a_jour.xlsx')
OUT = Path(__file__).resolve().parent.parent / 'data' / 'seed' / 'workshops.json'

def clean(s):
    if s is None: return ''
    s = str(s).replace('–', '-').replace('—', '-').replace('\xa0', ' ')
    return re.sub(r'[ \t]+', ' ', s).strip()

def bullets(s):
    return [clean(re.sub(r'^[•\-]\s*', '', l)) for l in str(s or '').split('\n') if clean(re.sub(r'^[•\-]\s*', '', l))]

def times(s):
    m = re.findall(r'(\d{2}):(\d{2})', s or '')
    return (f'{m[0][0]}:{m[0][1]}', f'{m[1][0]}:{m[1][1]}') if len(m) >= 2 else (None, None)

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb['Planning']
CODE = re.compile(r'^((J\d+)\.\d+|(C\d))\s+(.*)$', re.S)
blocks, block, session = [], None, None
for row in ws.iter_rows(min_row=4, values_only=True):
    a, b, c, d, e, f = (list(row) + [None] * 6)[:6]
    if all(x is None for x in row): continue
    A, B = clean(a), clean(b)
    if A == 'Horaire' or (block is None and not (b is None and c is None)): continue
    if b is None and c is None and a is not None:
        text = str(a)
        if text.startswith('Objectifs'):
            block['objectives'] = bullets(text.split(':', 1)[1])
        elif text.startswith(('Questionnaire', 'RACI')):
            for line in text.split('\n'):
                line = clean(line)
                if line.startswith('Questionnaire'): block['questionnaire'] = line.split(':', 1)[1].strip()
                elif line.startswith('RACI'): block['raci'].append(line.split('·', 1)[1].strip())
                elif line.startswith('Risques'): block['risks'] = [clean(r) for r in line.split(':', 1)[1].split('·')]
        elif text.startswith('Livrable'):
            block['deliverable'] = clean(text.split(':', 1)[1])
        elif text.startswith(('PV ', 'Un PV')):
            block['pv'] = clean(text)
        else:
            block = {'header': clean(text), 'objectives': [], 'questionnaire': '', 'raci': [], 'risks': [],
                     'deliverable': '', 'pv': '', 'sessions': [], 'agenda': []}
            blocks.append(block)
        continue
    if B.startswith('›'):
        start, end = times(A)
        m = re.search(r'\((\d+) min\)\s*$', B)
        title = clean(re.sub(r'\(\d+ min\)\s*$', '', B[1:]))
        session['steps'].append({'start': start, 'end': end, 'title': title, 'minutes': int(m.group(1)) if m else None})
        continue
    m = CODE.match(B)
    start, end = times(A)
    if m:
        title = m.group(4)
        added = 'Séquence ajoutée' in title
        title = clean(title.replace('Séquence ajoutée', ''))
        session = {'code': m.group(1), 'atelier': m.group(2) or m.group(3), 'start': start, 'end': end,
                   'allDay': 'Journée' in A, 'title': title, 'content': clean(c), 'owner': clean(d),
                   'outcome': clean(e), 'questions': bullets(f), 'added': added, 'steps': []}
        block['sessions'].append(session)
        block['agenda'].append({'kind': 'session', 'code': session['code']})
    elif B:
        block['agenda'].append({'kind': 'break', 'start': start, 'end': end, 'title': B})

# Dates et thèmes depuis l'onglet de suivi des PV
pv, prep = {}, {}
for row in wb['Suivi des PV'].iter_rows(min_row=4, values_only=True):
    k = clean(row[0])
    if re.fullmatch(r'J\d+|C\d', k or '') and k not in pv:
        fmt = lambda x: x.strftime('%Y-%m-%d') if hasattr(x, 'strftime') else None
        pv[k] = {'theme': clean(row[1]), 'date': fmt(row[2]), 'pvDue': fmt(row[3]), 'deliverableDue': fmt(row[4]),
                 'validationDue': fmt(row[5]), 'pvStatus': clean(row[6])}
    elif k and row[1] and str(row[1]).startswith('A') and hasattr(row[2], 'strftime'):
        for key in re.findall(r'J\d+', k): prep[key] = {'tab': clean(row[1]), 'due': row[2].strftime('%Y-%m-%d'), 'questions': row[3], 'status': clean(row[6])}

ateliers = []
for blk in blocks:
    if not blk['sessions']: continue
    codes = sorted({s['atelier'] for s in blk['sessions']}, key=lambda x: (x[0], int(x[1:])))
    if codes[0].startswith('C'):
        for s in blk['sessions']:
            ateliers.append({'id': s['atelier'], 'title': pv.get(s['atelier'], {}).get('theme') or s['title'],
                             'phase': 'Cadrage complémentaire V2', 'objectives': blk['objectives'] if s is blk['sessions'][0] else [],
                             'questionnaire': '', 'raci': blk['raci'], 'risks': [], 'deliverable': blk['deliverable'],
                             'pv': blk['pv'], 'sessions': [s], 'agenda': [{'kind': 'session', 'code': s['code']}],
                             **{k: v for k, v in pv.get(s['atelier'], {}).items() if k != 'theme'}})
        continue
    aid = codes[0]
    info = pv.get(aid, {})
    ateliers.append({'id': aid, 'title': info.get('theme') or blk['header'], 'phase': 'Ateliers de cadrage d\'Abidjan',
                     'objectives': blk['objectives'], 'questionnaire': blk['questionnaire'], 'raci': blk['raci'],
                     'risks': blk['risks'], 'deliverable': blk['deliverable'], 'pv': blk['pv'],
                     'sessions': blk['sessions'], 'agenda': blk['agenda'], 'prep': prep.get(aid),
                     **{k: v for k, v in info.items() if k != 'theme'}})

ateliers.sort(key=lambda w: (w['id'][0] == 'C', int(w['id'][1:])))
OUT.write_text(json.dumps({'source': SRC.name, 'workshops': ateliers}, ensure_ascii=False, indent=2) + '\n')
for w in ateliers:
    print(w['id'], w.get('date'), w['title'][:50], len(w['sessions']), 'sessions', sum(len(s['steps']) for s in w['sessions']), 'étapes')
